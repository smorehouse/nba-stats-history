#!/usr/bin/env python3
"""
Download the fantasy league Google Sheet and sync team/salary data into nba.sqlite.

Usage:
    uv run --with openpyxl python3 loader/sync_fantasy.py <GOOGLE_SHEET_ID> <DB_PATH>
"""

import re
import sqlite3
import sys
import unicodedata
from difflib import SequenceMatcher
from io import BytesIO
from urllib.request import urlopen

from openpyxl import load_workbook

TEAM_SHEETS = [
    "76ers", "Bucks", "Bulls", "Cavaliers", "Celtics", "Clippers",
    "Grizzlies", "Hawks", "Heat", "Hornets", "Jazz", "Kings",
    "Knicks", "Lakers", "Magic", "Mavericks", "Nets", "Nuggets",
    "Pacers", "Pelicans", "Pistons", "Raptors", "Rockets", "Spurs",
    "Suns", "Thunder", "Timberwolves", "Trailblazers", "Warriors", "Wizards",
]

SKIP_NAMES = {
    "total roster salary", "total paid out", "used salary", "salary cap",
    "available cap", "draft picks", "players", "",
}


def normalize_name(name: str) -> str:
    name = unicodedata.normalize("NFD", name)
    name = "".join(c for c in name if unicodedata.category(c) != "Mn")
    name = name.lower().strip()
    name = re.sub(r"\s+(jr\.?|sr\.?|ii|iii|iv)$", "", name)
    name = re.sub(r"[.'\-]", "", name)
    name = re.sub(r"\s+", " ", name)
    return name


def fuzzy_match(name: str, candidates: dict[str, int], threshold: float = 0.82) -> int | None:
    norm = normalize_name(name)
    best_score = 0.0
    best_id = None
    for cand_name, player_id in candidates.items():
        score = SequenceMatcher(None, norm, normalize_name(cand_name)).ratio()
        if score > best_score:
            best_score = score
            best_id = player_id
    if best_score >= threshold:
        return best_id
    return None


def parse_team_sheet(ws) -> list[dict]:
    """Parse a team worksheet, return list of {name, salary}."""
    players = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, values_only=True), 1):
        if row_idx <= 2:
            continue
        name = str(row[0]).strip() if row[0] else ""
        if not name or name.lower() in SKIP_NAMES:
            continue
        if name.lower().startswith(("total ", "used ", "available ", "salary ", "draft pick")):
            continue
        if re.match(r"^[\d.$]+$", name):
            continue

        salary = None
        if len(row) > 1 and row[1] is not None:
            try:
                salary = round(float(row[1]), 2)
            except (ValueError, TypeError):
                pass

        players.append({"name": name, "salary": salary, "is_dleague": False})

    # Also check column I (index 8) for d-league roster
    d_league_started = False
    for row in ws.iter_rows(min_row=1, values_only=True):
        if len(row) <= 8:
            continue
        val = str(row[8]).strip() if row[8] else ""
        if val.lower() == "d-league roster":
            d_league_started = True
            continue
        if d_league_started and val:
            if val.lower() in SKIP_NAMES or val.lower().startswith("waived"):
                continue
            clean = re.sub(r"\s*\(\$[\d.]+\)\s*$", "", val).strip()
            salary = None
            salary_match = re.search(r"\(\$([\d.]+)\)", val)
            if salary_match:
                salary = round(float(salary_match.group(1)), 2)
            elif len(row) > 9 and row[9] is not None:
                try:
                    salary = round(float(row[9]), 2)
                except (ValueError, TypeError):
                    pass
            if clean and not re.match(r"^[\d.$]+$", clean):
                players.append({"name": clean, "salary": salary, "is_dleague": True})

    return players


def main():
    if len(sys.argv) != 3:
        print(f"Usage: {sys.argv[0]} <GOOGLE_SHEET_ID> <DB_PATH>", file=sys.stderr)
        sys.exit(1)

    sheet_id = sys.argv[1]
    db_path = sys.argv[2]

    print("Downloading fantasy league spreadsheet...")
    url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
    data = urlopen(url).read()
    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)

    print("Parsing team rosters...")
    all_fantasy = []
    for team_name in TEAM_SHEETS:
        if team_name not in wb.sheetnames:
            print(f"  Warning: sheet '{team_name}' not found, skipping")
            continue
        ws = wb[team_name]
        players = parse_team_sheet(ws)
        for p in players:
            all_fantasy.append({**p, "team": team_name})
        print(f"  {team_name}: {len(players)} players")

    wb.close()

    conn = sqlite3.connect(db_path)
    conn.execute("DROP TABLE IF EXISTS fantasy_rosters")
    conn.execute("CREATE TABLE fantasy_rosters ("
                 "player_id INTEGER PRIMARY KEY, "
                 "fantasy_team TEXT NOT NULL, "
                 "salary REAL, "
                 "is_dleague INTEGER NOT NULL DEFAULT 0)")

    nba_players = {row[1]: row[0] for row in conn.execute("SELECT player_id, player_name FROM players")}
    print(f"\nMatching {len(all_fantasy)} fantasy players to {len(nba_players)} NBA players...")

    matched = 0
    unmatched = []
    for fp in all_fantasy:
        player_id = fuzzy_match(fp["name"], nba_players)
        if player_id:
            conn.execute(
                "INSERT OR REPLACE INTO fantasy_rosters (player_id, fantasy_team, salary, is_dleague) VALUES (?, ?, ?, ?)",
                (player_id, fp["team"], fp["salary"], 1 if fp["is_dleague"] else 0),
            )
            matched += 1
        else:
            unmatched.append(fp["name"])

    conn.commit()
    conn.close()

    print(f"Matched: {matched}/{len(all_fantasy)}")
    if unmatched:
        print(f"Unmatched ({len(unmatched)}): {', '.join(unmatched[:10])}")
        if len(unmatched) > 10:
            print(f"  ... and {len(unmatched) - 10} more")


if __name__ == "__main__":
    main()
